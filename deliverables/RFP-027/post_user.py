"""Post-process Alaa's edited Rev 02 (user_edit.xlsx): neutral wording, bullet notes, wrapping, no same-tab links on the Assessment totals."""
import openpyxl, re, copy
from openpyxl.styles import Alignment
u = openpyxl.load_workbook('user_edit.xlsx')
PH = [
 ("Rate: assessed allowance (Riyadh market, Sep-2026) - needs confirmation by the Contractor", "Rate: assessed allowance (Riyadh market, Sep-2026)"),
 ("Rate: assessed allowance (Riyadh market, Sep-2026) - needs confirmation", "Rate: assessed allowance (Riyadh market, Sep-2026)"),
 (" - needs confirmation by the Contractor's method statement", "; the Contractor's method statement is outstanding"),
 (", needs confirmation by the Contractor's method statement", "; the Contractor's method statement is outstanding"),
 (" - needs confirmation by the Contractor", ""), (", needs confirmation by the Contractor", ""),
 ("The rate needs confirmation by the Contractor", "The rate is an assessed allowance"),
 (" - needs confirmation of the water source", "; the water source is unconfirmed"),
 ("'Needs confirmation' marks an assessed allowance the Contractor has not yet substantiated; ", ""),
 ("- assumption pending the Contractor's substantiation", "- assessed allowance"),
 ("assumption pending the Contractor's substantiation", "assessed allowance"),
 ("an assumption pending the Contractor's plant schedule", "an assessed assumption; the Contractor's plant schedule is outstanding"),
 ("an assessed assumption pending the Contractor's method statement", "an assessed assumption; the Contractor's method statement is outstanding"),
]
SENT = re.compile(r"(?<=[a-z0-9\)\]'])\. (?=[A-Z(])")
def bulletise(text):
    parts = [p for p in SENT.split(text) if p]
    out = []
    for p in parts:
        if out and (out[-1].endswith(' No') or out[-1].endswith(' Sub-Clause') or len(p) < 25): out[-1] += '. ' + p
        else: out.append(p)
    if len(out) < 2: return text
    return '\n'.join('• ' + p.rstrip('.') for p in out)
BULLET_COLS = {'Build-Up': (7,), 'Programme': (9, 10), 'Build-Up Comparison': (8,), 'Assessment': (12,)}
n_ph = n_bul = 0
for ws in u:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith('='): continue
            new = v
            for a, b in PH: new = new.replace(a, b)
            if new != v: c.value = new; n_ph += 1; v = new
            if ws.title in BULLET_COLS and c.column in BULLET_COLS[ws.title] and c.row >= 6 and len(v) > 110 and '•' not in v:
                nv = bulletise(v)
                if nv != v:
                    c.value = nv; n_bul += 1
                    al = c.alignment; c.alignment = Alignment(horizontal=al.horizontal or 'left', vertical=al.vertical or 'center', wrap_text=True)
# wrap the Build-Up mirrors on the Programme tab
pg = u['Programme']
for row in pg.iter_rows(min_col=2, max_col=3):
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("='Build-Up'!"):
            al = c.alignment; c.alignment = Alignment(horizontal=al.horizontal or 'left', vertical='center', wrap_text=True)
# no same-tab links on the Assessment arithmetic
asm = u['Assessment']; n_rm = 0
for row in asm.iter_rows():
    for c in row:
        if c.hyperlink and (c.hyperlink.location or '').startswith(("'Assessment'", "Assessment")):
            c.hyperlink = None; n_rm += 1
            f = copy.copy(c.font); c.font = openpyxl.styles.Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color='000000', u=None)
# row heights for bulleted cells
import math
from openpyxl.utils import get_column_letter
for name in BULLET_COLS:
    ws = u[name]
    merged = {}
    for rg in ws.merged_cells.ranges:
        c1, r1, c2, r2 = rg.bounds
        if r1 == r2: merged[(r1, c1)] = sum((ws.column_dimensions[get_column_letter(c)].width or 8.66) for c in range(c1, c2 + 1))
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or c.value.startswith('=') or not (c.alignment and c.alignment.wrap_text): continue
            width = merged.get((c.row, c.column), ws.column_dimensions[get_column_letter(c.column)].width or 8.66)
            sz = c.font.sz or 10; cpl = max(8, width * 1.15 * 10 / sz)
            need = sum(max(1, math.ceil(len(p) / cpl)) for p in c.value.split('\n')) * (sz * 1.28) + 4
            cur = ws.row_dimensions[c.row].height
            if cur is None or need > cur: ws.row_dimensions[c.row].height = round(need, 1)
u['Assessment'].page_setup.fitToHeight = 1
u.save('stageU.xlsx')
print('phrases cleaned', n_ph, '| bulleted', n_bul, '| assessment same-tab links removed', n_rm)
