"""Semantic hyperlink audit for an assessment workbook.
For every link: the target row must carry a label; a lookup formula (MATCH("<activity ID>")) must link to that activity's
row on the source tab; a SUMPRODUCT count must list the linked activity in its note column; any other formula must link to
its first referenced source (another tab, or another row on the same tab; ranges and fixed $-constants ignored).
Also lists formulas with a cross-reference but no link. A link on a multi-input formula is navigation to one source only.

usage: python audit_links.py workbook.xlsx [--source-tab "XER WBS"] [--id-prefix QCD18TSE] [--note-col 10]
"""
import openpyxl, re, sys, argparse
ap = argparse.ArgumentParser(); ap.add_argument('path'); ap.add_argument('--source-tab', default='XER WBS'); ap.add_argument('--id-prefix', default='QCD18TSE'); ap.add_argument('--note-col', type=int, default=10)
a = ap.parse_args()
f = openpyxl.load_workbook(a.path)
REF = re.compile(r"(?:'([^']+)'!|(?<![A-Za-z_])([A-Za-z][A-Za-z ]*?)!)?\$?([A-Z]{1,2})\$?(\d+)(?![\d(])")
xid = {}
if a.source_tab in f.sheetnames:
    x = f[a.source_tab]
    for row in x.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if isinstance(c.value, str) and c.value.startswith(a.id_prefix) and c.value not in xid: xid[c.value] = c.row   # first occurrence: the activity table, not the relationships list
def tgt(loc):
    m = re.match(r"'?([^'!]+)'?!([A-Z]+)(\d+)", loc or ''); return (m.group(1), m.group(2), int(m.group(3))) if m else (None, None, None)
def src(v, own, own_row):
    for m in REF.finditer(v):
        if v[m.end():m.end()+1] == ':' or v[m.start()-1:m.start()] == ':': continue
        sh = m.group(1) or m.group(2) or own; r = int(m.group(4))
        if sh not in f.sheetnames or (sh == own and r == own_row) or (sh == own and m.group(0).count('$') == 2): continue
        return sh, r
    return None
issues, missing, n = [], [], 0
for ws in f:
    for row in ws.iter_rows():
        for c in row:
            v = c.value if isinstance(c.value, str) else ''
            if c.hyperlink:
                n += 1; sh, col, r = tgt(c.hyperlink.location)
                if c.hyperlink.target: issues.append((ws.title, c.coordinate, 'external target', c.hyperlink.target)); continue
                if sh not in f.sheetnames or (f[sh].cell(r, 1).value in (None, '') and f[sh].cell(r, 2).value in (None, '')): issues.append((ws.title, c.coordinate, 'unlabelled or missing target', c.hyperlink.location)); continue
                ids = re.findall(r'MATCH\("(' + a.id_prefix + r'[A-Z0-9]+)"', v)
                if ids:
                    if sh != a.source_tab or xid.get(ids[0]) != r: issues.append((ws.title, c.coordinate, 'lookup id', ids[0], c.hyperlink.location))
                elif v.startswith('=SUMPRODUCT'):
                    note = ws.cell(c.row, a.note_col).value or ''
                    if sh == a.source_tab and str(f[sh].cell(r, 1).value) not in note: issues.append((ws.title, c.coordinate, 'counted id not in note', c.hyperlink.location))
                elif v.startswith('=') and not v.startswith('="'):
                    s = src(v, ws.title, c.row)
                    if s and s != (sh, r): issues.append((ws.title, c.coordinate, 'first source', v[:50], c.hyperlink.location))
            elif v.startswith('=') and not v.startswith('="') and ws.title not in (a.source_tab, 'Data', 'Project Info', 'Assessment'):
                if src(v, ws.title, c.row): missing.append((ws.title, c.coordinate, v[:50]))
print(f'links {n} | mislinked {len(issues)} | cross-references without a link {len(missing)}')
for i in issues[:40]: print('  MISLINK', i)
for m in missing[:40]: print('  UNLINKED', m)
sys.exit(1 if issues else 0)
