"""Turn essay notes into bullets, wrap mirrored text, remove invitations to the Contractor, drop same-tab links on the
Assessment totals, and recalculate row heights. Run on a hand-edited delivery (openpyxl drops cover drawings: follow with
restore_drawings.py) or fold the same functions into a build script.

usage: python bulletise_notes.py in.xlsx out.xlsx
"""
import openpyxl, re, math, sys, copy
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
src, dst = sys.argv[1], sys.argv[2]
u = openpyxl.load_workbook(src)
B = '• '
PHRASES = [
 ("Rate: assessed allowance (Riyadh market, Sep-2026) - needs confirmation by the Contractor", "Rate: assessed allowance (Riyadh market, Sep-2026)"),
 (" - needs confirmation by the Contractor's method statement", "; the Contractor's method statement is outstanding"),
 (", needs confirmation by the Contractor's method statement", "; the Contractor's method statement is outstanding"),
 (" - needs confirmation by the Contractor", ""), (", needs confirmation by the Contractor", ""), (" - needs confirmation", ""),
 ("an assumption pending the Contractor's plant schedule", "an assessed assumption; the Contractor's plant schedule is outstanding"),
 ("an assessed assumption pending the Contractor's method statement", "an assessed assumption; the Contractor's method statement is outstanding"),
 ("- assumption pending the Contractor's substantiation", "- assessed allowance"), ("assumption pending the Contractor's substantiation", "assessed allowance"),
 ("'Needs confirmation' marks an assessed allowance the Contractor has not yet substantiated; ", ""),
]
SENT = re.compile(r"(?<=[a-z0-9\)\]'])\. (?=[A-Z(])")
def bulletise(text):
    head = ''
    m = re.match(r'^(\d\. )', text)
    if m: head = m.group(1); text = text[len(head):]
    parts = [p for p in SENT.split(text) if p]; out = []
    for p in parts:
        if out and (out[-1].endswith(' No') or out[-1].endswith(' Sub-Clause') or len(p) < 25): out[-1] += '. ' + p
        else: out.append(p)
    if len(out) < 2: return head + text
    return head + out[0].rstrip('.') + '\n' + '\n'.join(B + p.rstrip('.') for p in out[1:])
NOTE_COLS = {'Build-Up': (1, 7), 'Programme': (1, 9, 10), 'Build-Up Comparison': (1, 8), 'Assessment': (1, 12)}
for ws in u:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith('='): continue
            for p, q in PHRASES: v = v.replace(p, q)
            if ws.title in NOTE_COLS and c.column in NOTE_COLS[ws.title] and c.row >= 4 and len(v) > 110 and B not in v: v = bulletise(v)
            if v != c.value:
                c.value = v; al = c.alignment; c.alignment = Alignment(horizontal=al.horizontal or 'left', vertical=al.vertical or 'center', wrap_text=True)
if 'Programme' in u.sheetnames:
    for row in u['Programme'].iter_rows(min_col=2, max_col=3):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("='Build-Up'!"):
                al = c.alignment; c.alignment = Alignment(horizontal=al.horizontal or 'left', vertical='center', wrap_text=True)
if 'Assessment' in u.sheetnames:
    for row in u['Assessment'].iter_rows():
        for c in row:
            if c.hyperlink and (c.hyperlink.location or '').lstrip("'").startswith('Assessment'):
                c.hyperlink = None; f = copy.copy(c.font); c.font = Font(name=f.name, sz=f.sz, b=f.b, i=f.i, color='000000')
for name in NOTE_COLS:
    if name not in u.sheetnames: continue
    ws = u[name]; merged = {}
    for rg in ws.merged_cells.ranges:
        c1, r1, c2, r2 = rg.bounds
        if r1 == r2: merged[(r1, c1)] = sum((ws.column_dimensions[get_column_letter(cc)].width or 8.66) for cc in range(c1, c2 + 1))
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or c.value.startswith('=') or not (c.alignment and c.alignment.wrap_text): continue
            width = merged.get((c.row, c.column), ws.column_dimensions[get_column_letter(c.column)].width or 8.66)
            sz = c.font.sz or 10; cpl = max(8, width * 1.15 * 10 / sz)
            need = sum(max(1, math.ceil(len(p) / cpl)) for p in c.value.split('\n')) * (sz * 1.28) + 4
            cur = ws.row_dimensions[c.row].height
            if cur is None or need > cur: ws.row_dimensions[c.row].height = round(need, 1)
u.save(dst); print('written', dst)
